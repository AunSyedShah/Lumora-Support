"""
Turn an uploaded file into a flat list of text Blocks (paragraphs / headings) with page numbers.

Each parser only extracts text; deciding what is a section heading happens in chunking.py.
"""

import io
import re
import unicodedata
from dataclasses import dataclass

import docx
import pymupdf


@dataclass
class Block:
    text: str
    page: int | None = None  # 1-based page number, PDFs only
    heading_hint: bool = False  # bold PDF text, a DOCX "Heading" style, or a Markdown '#'


@dataclass
class ParsedDocument:
    blocks: list[Block]
    page_count: int | None = None

    @property
    def full_text(self):
        return "\n".join(b.text for b in self.blocks)


def _clean(text):
    # NFKC turns PDF ligatures like "ﬀ" (one glyph) back into plain "ff", fancy spaces into spaces, etc.
    text = unicodedata.normalize("NFKC", text)
    return re.sub(r"\s+", " ", text).strip()


def parse_pdf(data: bytes) -> ParsedDocument:
    blocks = []
    with pymupdf.open(stream=data, filetype="pdf") as pdf:
        for page_number, page in enumerate(pdf, start=1):
            for block in page.get_text("dict")["blocks"]:
                spans = [s for line in block.get("lines", []) for s in line["spans"] if s["text"].strip()]
                if not spans:
                    continue  # image block or blank
                text = _clean(" ".join(s["text"] for s in spans))
                is_bold = all(s["flags"] & pymupdf.TEXT_FONT_BOLD or "Bold" in s["font"] for s in spans)
                blocks.append(Block(text=text, page=page_number, heading_hint=is_bold))
        return ParsedDocument(blocks=blocks, page_count=pdf.page_count)


def parse_docx(data: bytes) -> ParsedDocument:
    document = docx.Document(io.BytesIO(data))
    blocks = []
    for paragraph in document.paragraphs:
        text = _clean(paragraph.text)
        if not text:
            continue
        style = (paragraph.style.name or "").lower() if paragraph.style is not None else ""
        runs = [r for r in paragraph.runs if r.text.strip()]
        all_bold = bool(runs) and all(r.bold for r in runs)
        blocks.append(Block(text=text, heading_hint=style.startswith(("heading", "title")) or all_bold))

    # Tables (e.g. SLA or routing tables): one block per row, cells joined with " | ".
    for table in document.tables:
        for row in table.rows:
            text = " | ".join(_clean(cell.text) for cell in row.cells if cell.text.strip())
            if text:
                blocks.append(Block(text=text))
    return ParsedDocument(blocks=blocks)


def parse_text(data: bytes) -> ParsedDocument:
    """Plain text and Markdown: one block per non-empty line; Markdown '#' lines are headings."""
    blocks = []
    for line in data.decode("utf-8", errors="replace").splitlines():
        stripped = line.strip()
        if not stripped:
            continue
        is_md_heading = stripped.startswith("#")
        blocks.append(Block(text=_clean(stripped.lstrip("#")), heading_hint=is_md_heading))
    return ParsedDocument(blocks=blocks)


PARSERS = {
    ".pdf": parse_pdf,
    ".docx": parse_docx,
    ".txt": parse_text,
    ".md": parse_text,
}


def parse_file(extension: str, data: bytes) -> ParsedDocument:
    return PARSERS[extension](data)


# ---------------------------------------------------------------------------
# Metadata header, e.g. at the top of a document:
#   Document ID: REF-POL | Version: 2.0 | Type: Policy | Status: Active
#   Effective Date: 2026-01-01 | Expiry Date: 2026-12-31 | Category: REFUND
# ---------------------------------------------------------------------------

HEADER_PATTERNS = {
    "doc_id": r"Document\s*ID\s*[:\-]\s*([A-Za-z0-9_\-]+)",
    "version": r"Version\s*[:\-]\s*v?([0-9][0-9A-Za-z.\-]*)",
    "doc_type": r"(?:Document\s*)?Type\s*[:\-]\s*([A-Za-z ]+?)\s*(?:\||$)",
    "status": r"Status\s*[:\-]\s*([A-Za-z]+)",
    "effective_date": r"Effective(?:\s*Date)?\s*[:\-]\s*(\d{4}-\d{2}-\d{2})",
    "expiry_date": r"Expir(?:y|es)(?:\s*Date)?\s*[:\-]\s*(\d{4}-\d{2}-\d{2})",
    "category": r"Category\s*[:\-]\s*([A-Z0-9_]+)",
}


METADATA_LINE = re.compile(r"Document\s*ID|Version\s*:|Effective|Status\s*:", re.IGNORECASE)
SECTION_START = re.compile(r"^\d{1,2}(\.\d{1,2})*\.?\s")  # "2.1 Refunds" is a section, never a title


def extract_header_metadata(parsed: ParsedDocument) -> dict:
    """Look for metadata fields in the first few blocks. Returns only what was found."""
    header_text = "\n".join(b.text for b in parsed.blocks[:8])
    found = {}
    for field, pattern in HEADER_PATTERNS.items():
        match = re.search(pattern, header_text, flags=re.IGNORECASE | re.MULTILINE)
        if match:
            found[field] = match.group(1).strip()

    # The title is the first block that is not a metadata line. If a numbered section comes
    # first, the document has no title line (the caller then falls back to the file name).
    for block in parsed.blocks[:3]:
        if SECTION_START.match(block.text):
            break
        if not METADATA_LINE.search(block.text):
            found["title"] = re.sub(r"^title\s*:\s*", "", block.text, flags=re.IGNORECASE)[:200]
            break
    return found


def count_header_blocks(parsed: ParsedDocument, title: str | None) -> int:
    """How many leading blocks are just the title / metadata header (not policy content)."""
    count = 0
    for block in parsed.blocks[:6]:
        if block.text == title or METADATA_LINE.search(block.text):
            count += 1
        else:
            break
    return count
