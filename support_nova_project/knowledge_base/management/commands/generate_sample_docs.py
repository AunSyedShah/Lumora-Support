"""
Render the Markdown sources in sample_documents/source/ into real PDF / DOCX files.

The target format comes from the file name:  refund_policy_v2.pdf.md -> refund_policy_v2.pdf
Supported Markdown: '# title', '## 1. heading', '### 1.1 sub-heading', '- bullet', plain paragraphs.

Usage:  python manage.py generate_sample_docs
"""

import html
from pathlib import Path

import docx
import pymupdf
from django.conf import settings
from django.core.management.base import BaseCommand

SOURCE_DIR = settings.BASE_DIR / "sample_documents" / "source"
OUTPUT_DIR = settings.BASE_DIR / "sample_documents"


def parse_markdown(text):
    """Yield (kind, text) pairs: kind is title / h1 / h2 / bullet / para."""
    for line in text.splitlines():
        line = line.strip()
        if not line:
            continue
        if line.startswith("### "):
            yield "h2", line[4:]
        elif line.startswith("## "):
            yield "h1", line[3:]
        elif line.startswith("# "):
            yield "title", line[2:]
        elif line.startswith("- "):
            yield "bullet", line[2:]
        else:
            yield "para", line


def render_docx(items, path):
    document = docx.Document()
    for kind, text in items:
        if kind == "title":
            document.add_heading(text, level=0)
        elif kind == "h1":
            document.add_heading(text, level=1)
        elif kind == "h2":
            document.add_heading(text, level=2)
        elif kind == "bullet":
            document.add_paragraph(text, style="List Bullet")
        else:
            document.add_paragraph(text)
    document.save(path)


def render_pdf(items, path):
    tags = {"title": "h1", "h1": "h2", "h2": "h3", "para": "p"}
    parts = []
    for kind, text in items:
        text = html.escape(text)
        parts.append(f"<ul><li>{text}</li></ul>" if kind == "bullet" else f"<{tags[kind]}>{text}</{tags[kind]}>")

    story = pymupdf.Story(html="\n".join(parts), user_css="body {font-family: sans-serif; font-size: 11pt;}")
    writer = pymupdf.DocumentWriter(str(path))
    page_rect = pymupdf.paper_rect("a4")
    content_rect = page_rect + (50, 50, -50, -50)  # 50pt margins
    more = True
    while more:
        device = writer.begin_page(page_rect)
        more, _ = story.place(content_rect)
        story.draw(device)
        writer.end_page()
    writer.close()


class Command(BaseCommand):
    help = "Generate sample PDF/DOCX knowledge-base documents from Markdown sources."

    def handle(self, *args, **options):
        renderers = {".pdf": render_pdf, ".docx": render_docx}
        for source in sorted(SOURCE_DIR.glob("*.md")):
            target = OUTPUT_DIR / source.stem  # "refund_policy_v2.pdf.md" -> "refund_policy_v2.pdf"
            renderer = renderers.get(Path(source.stem).suffix)
            if renderer is None:
                self.stderr.write(f"Skipped {source.name}: name must end in .pdf.md or .docx.md")
                continue
            renderer(list(parse_markdown(source.read_text(encoding="utf-8"))), target)
            self.stdout.write(f"Created {target.relative_to(settings.BASE_DIR)}")
